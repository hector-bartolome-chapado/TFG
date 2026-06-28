from __future__ import annotations

import argparse
import json
import pathlib
import re
import time
import unicodedata
from typing import Any

from openpyxl import load_workbook

from scripts.simple_extractor.config import (
    DEFAULT_EMBED_MODEL,
    DEFAULT_LLAMUS_BASE_URL,
    get_api_key,
)
from scripts.simple_extractor.embeddings import request_embedding


DEFAULT_ROWS_PER_CHUNK = 12


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9]+", "_", ascii_value).strip("_")
    return slug or "documento"


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def write_jsonl(path: pathlib.Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def row_to_text(row_number: int, values: tuple[Any, ...]) -> str:
    cells = [clean_cell(value) for value in values]
    cells = [cell for cell in cells if cell]
    if not cells:
        return ""
    return f"Fila {row_number}: " + " | ".join(cells)


def build_blocks_from_xlsx(path: pathlib.Path, doc_id: str, rows_per_chunk: int) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[dict[str, Any]] = []

    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        pending_rows: list[tuple[int, str]] = []

        def flush() -> None:
            if not pending_rows:
                return
            row_start = pending_rows[0][0]
            row_end = pending_rows[-1][0]
            block_index = len(blocks) + 1
            text = "\n".join(row_text for _, row_text in pending_rows)
            blocks.append(
                {
                    "block_id": f"{doc_id}::sheet{sheet_index:02d}::rows{row_start}-{row_end}::b{block_index}",
                    "doc_id": doc_id,
                    "source_file": str(path),
                    "sheet": worksheet.title,
                    "sheet_index": sheet_index,
                    "row_start": row_start,
                    "row_end": row_end,
                    "text": (
                        f"Documento: {path.name}\n"
                        f"Hoja: {worksheet.title}\n"
                        f"Rango de filas: {row_start}-{row_end}\n"
                        f"{text}"
                    ),
                }
            )
            pending_rows.clear()

        for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            text = row_to_text(row_number, values)
            if not text:
                continue
            pending_rows.append((row_number, text))
            if len(pending_rows) >= rows_per_chunk:
                flush()
        flush()

    return blocks


def build_chunks(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    for block in blocks:
        parent_id = f"{block['block_id']}::parent"
        chunk_id = f"{parent_id}::child::1"
        chunks.append(
            {
                "chunk_id": chunk_id,
                "parent_id": parent_id,
                "doc_id": block["doc_id"],
                "source_file": block["source_file"],
                "sheet": block["sheet"],
                "sheet_index": block["sheet_index"],
                "row_start": block["row_start"],
                "row_end": block["row_end"],
                "text": block["text"],
                "retrieval_text": block["text"],
                "parent_text": block["text"],
                "chunk_index": 1,
                "chunk_count": 1,
                "chunking_strategy": "xlsx_rows_parent_child",
                "rows_per_chunk": block["row_end"] - block["row_start"] + 1,
            }
        )
    return chunks


def embed_chunks(
    chunks: list[dict[str, Any]],
    model: str,
    base_url: str,
    api_key: str,
) -> list[dict[str, Any]]:
    embedded: list[dict[str, Any]] = []
    total = len(chunks)
    for index, chunk in enumerate(chunks, start=1):
        print(f"[{index}/{total}] embedding {chunk['chunk_id']}", flush=True)
        row = dict(chunk)
        row["model"] = model
        row["embedding"] = request_embedding(
            chunk["retrieval_text"],
            model=model,
            base_url=base_url,
            api_key=api_key,
        )
        embedded.append(row)
    return embedded


def process_xlsx(
    path: pathlib.Path,
    output_root: pathlib.Path,
    rows_per_chunk: int,
    embed: bool,
    model: str,
    base_url: str,
    api_key: str | None,
) -> dict[str, Any]:
    started = time.perf_counter()
    doc_id = slugify(path.stem)
    blocks = build_blocks_from_xlsx(path, doc_id=doc_id, rows_per_chunk=rows_per_chunk)
    chunks = build_chunks(blocks)

    blocks_path = output_root / "blocks" / f"{doc_id}.jsonl"
    chunks_path = output_root / "chunks" / f"{doc_id}.jsonl"
    write_jsonl(blocks_path, blocks)
    write_jsonl(chunks_path, chunks)

    result: dict[str, Any] = {
        "doc_id": doc_id,
        "source_file": str(path),
        "block_count": len(blocks),
        "chunk_count": len(chunks),
        "rows_per_chunk": rows_per_chunk,
        "blocks_path": str(blocks_path),
        "chunks_path": str(chunks_path),
    }

    if embed:
        if not api_key:
            raise ValueError("No hay API key para embeddings. Usa TFG/.llamus_api_key o LLAMUS_API_KEY.")
        embeddings_path = output_root / "embeddings" / f"{doc_id}.jsonl"
        embedded = embed_chunks(chunks, model=model, base_url=base_url, api_key=api_key)
        write_jsonl(embeddings_path, embedded)
        result["embeddings_path"] = str(embeddings_path)
        result["embedding_model"] = model
        result["embedding_dimension"] = len(embedded[0]["embedding"]) if embedded else 0

    result["elapsed_seconds"] = time.perf_counter() - started
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Procesa Excel de PGE y genera JSONL/embeddings.")
    parser.add_argument("paths", nargs="+", type=pathlib.Path)
    parser.add_argument("--output-root", type=pathlib.Path, default=pathlib.Path("RESULTADOS EMBEDDING"))
    parser.add_argument("--rows-per-chunk", type=int, default=DEFAULT_ROWS_PER_CHUNK)
    parser.add_argument("--embed", action="store_true")
    parser.add_argument("--model", default=DEFAULT_EMBED_MODEL)
    parser.add_argument("--base-url", default=DEFAULT_LLAMUS_BASE_URL)
    args = parser.parse_args()

    project_root = pathlib.Path(__file__).resolve().parents[2]
    api_key = get_api_key(project_root)
    summaries = []
    for path in args.paths:
        summaries.append(
            process_xlsx(
                path=path,
                output_root=args.output_root,
                rows_per_chunk=args.rows_per_chunk,
                embed=args.embed,
                model=args.model,
                base_url=args.base_url,
                api_key=api_key,
            )
        )

    summary_path = args.output_root / "ingesta_eval" / "pge_xlsx_embeddings_summary.json"
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summaries, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"summary_path": str(summary_path), "documents": summaries}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
